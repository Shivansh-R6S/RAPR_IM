import os
import logging
import networkx as nx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

SEED_METHODS = ["RPaR-IM", "CELF Greedy", "PageRank", "Degree Centrality", "Random"]
METRICS = ["Avg σ+", "Polarization", "Net Influence", "Seed Efficiency"]
METRIC_KEYS = ["avg_sigma_plus", "polarization_index", "net_influence", "seed_efficiency"]
DIFFUSION_MODELS = ["PaIC-DGBC", "Signed IC", "Signed LT"]


def build_results_table(G: nx.DiGraph, seed_sizes: list,
                         seed_sets: dict, diffusion_fns: dict,
                         mc_runs: int = 10) -> dict:
    """
    Runs all evaluations and returns a nested dict:
        results[diffusion_model][seed_method][k] = {metric: value, ...}

    seed_sets: {method_name: {k: [node, ...]}}
    diffusion_fns: {model_name: callable}
    """
    from src.baselines import evaluate
    results = {model: {method: {} for method in SEED_METHODS} for model in DIFFUSION_MODELS}

    for model_name, diff_fn in diffusion_fns.items():
        for method in SEED_METHODS:
            for k in seed_sizes:
                seeds = seed_sets[method].get(k, [])
                logger.info(f"Evaluating {model_name} | {method} | k={k}...")
                results[model_name][method][k] = evaluate(G, seeds, diff_fn, mc_runs=mc_runs)

    return results


def write_excel(results: dict, seed_sizes: list, output_path: str):
    """
    Writes results to a formatted Excel workbook.
    One sheet per diffusion model. Each sheet has:
    - Rows: seed sizes (k=5,10,15,20,25), each size takes 4 rows (one per metric)
    - Columns: seed selection methods, each method takes 4 sub-columns (one per metric)
    """
    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    # styles
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    subheader_font = Font(name="Arial", bold=True, size=9)
    cell_font = Font(name="Arial", size=9)
    method_fill = PatternFill("solid", start_color="1F4E79")   # dark blue
    metric_fill = PatternFill("solid", start_color="2E75B6")   # mid blue
    k_fill = PatternFill("solid", start_color="D6E4F0")        # light blue
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for model_name in DIFFUSION_MODELS:
        ws = wb.create_sheet(title=model_name)

        # --- build column layout ---
        # col 1: "k" label
        # then for each method: 4 metric sub-columns
        n_methods = len(SEED_METHODS)
        n_metrics = len(METRICS)

        # row 1: method headers (merged across 4 metric cols each)
        # row 2: metric sub-headers
        # row 3+: data (one row per seed size per metric)

        ws.column_dimensions["A"].width = 14

        # method headers - row 1
        ws.cell(row=1, column=1, value="Seed Size").font = subheader_font
        ws.cell(row=1, column=1).fill = PatternFill("solid", start_color="1F4E79")
        ws.cell(row=1, column=1).font = header_font
        ws.cell(row=1, column=1).alignment = center
        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 30

        for m_idx, method in enumerate(SEED_METHODS):
            start_col = 2 + m_idx * n_metrics
            end_col = start_col + n_metrics - 1

            ws.merge_cells(
                start_row=1, start_column=start_col,
                end_row=1, end_column=end_col
            )
            cell = ws.cell(row=1, column=start_col, value=method)
            cell.font = header_font
            cell.fill = method_fill
            cell.alignment = center
            cell.border = border

            # metric sub-headers - row 2
            for met_idx, metric in enumerate(METRICS):
                col = start_col + met_idx
                c = ws.cell(row=2, column=col, value=metric)
                c.font = Font(name="Arial", bold=True, size=8, color="FFFFFF")
                c.fill = metric_fill
                c.alignment = center
                c.border = border
                ws.column_dimensions[get_column_letter(col)].width = 13

        # data rows
        row = 3
        for k in seed_sizes:
            # k label spanning all metric rows for this seed size
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row + n_metrics - 1, end_column=1
            )
            k_cell = ws.cell(row=row, column=1, value=f"k = {k}")
            k_cell.font = Font(name="Arial", bold=True, size=9)
            k_cell.fill = k_fill
            k_cell.alignment = center
            k_cell.border = border

            for met_idx, metric_key in enumerate(METRIC_KEYS):
                data_row = row + met_idx
                ws.row_dimensions[data_row].height = 16

                for m_idx, method in enumerate(SEED_METHODS):
                    col = 2 + m_idx * n_metrics + met_idx
                    val = results[model_name][method][k].get(metric_key, "")
                    c = ws.cell(row=data_row, column=col, value=val)
                    c.font = cell_font
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border = border

                    # shade alternating method blocks lightly
                    if m_idx % 2 == 0:
                        c.fill = PatternFill("solid", start_color="F5F9FF")

            row += n_metrics

        # freeze panes below headers and right of k column
        ws.freeze_panes = "B3"

    wb.save(output_path)
    logger.info(f"Results saved to {output_path}")